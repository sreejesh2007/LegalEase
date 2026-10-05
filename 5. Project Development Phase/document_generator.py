from io import BytesIO
from html import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


def build_term_rows(data):
    return [
        ["Document Type", data.document_type],
        ["Party A", data.party_a],
        ["Party B", data.party_b],
        ["Effective Date", data.effective_date],
        ["Term", data.term],
        ["Jurisdiction", data.jurisdiction],
        ["Purpose / Role / Property", data.purpose],
        ["Payment / Consideration", data.consideration],
        ["Special Terms", data.special_terms or "Not specified"],
    ]


def build_text_document(data, draft):
    lines = []

    lines.append(data.logo_text or "LegalEase")
    lines.append("=" * 60)
    lines.append(data.document_type)
    lines.append("=" * 60)
    lines.append("")

    lines.append("KEY TERMS")
    lines.append("-" * 60)

    for label, value in build_term_rows(data):
        lines.append(f"{label}: {value}")

    lines.append("")
    lines.append("GENERATED DOCUMENT")
    lines.append("-" * 60)
    lines.append("")
    lines.append(draft)

    lines.append("")
    lines.append("-" * 60)
    lines.append(
        "LEGAL REVIEW NOTICE: This document is an AI-generated draft "
        "and should be reviewed by a qualified legal professional "
        "before actual use."
    )

    return "\n".join(lines)


def create_txt_file(data, draft):
    content = build_text_document(data, draft)

    return BytesIO(content.encode("utf-8"))


def create_docx_file(data, draft):
    document = Document()

    title = document.add_heading(data.logo_text or "LegalEase", 0)
    title.alignment = 1

    document.add_heading(data.document_type, level=1)

    document.add_heading("Key Terms", level=2)

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.style = "Table Grid"

    header = table.rows[0].cells
    header[0].text = "Term"
    header[1].text = "Details"

    for label, value in build_term_rows(data):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_heading("Generated Document", level=2)

    for line in draft.splitlines():
        paragraph = document.add_paragraph()

        run = paragraph.add_run(line)
        run.font.name = data.font
        run.font.size = Pt(11)

    document.add_paragraph("")
    notice = document.add_paragraph()

    run = notice.add_run(
        "LEGAL REVIEW NOTICE: This document is an AI-generated draft "
        "and should be reviewed by a qualified legal professional "
        "before actual use."
    )

    run.bold = True
    run.font.size = Pt(9)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output


def create_pdf_file(data, draft):
    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "LegalEaseTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        spaceAfter=10,
    )

    heading_style = ParagraphStyle(
        "LegalEaseHeading",
        parent=styles["Heading2"],
        fontSize=13,
        spaceBefore=10,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "LegalEaseBody",
        parent=styles["BodyText"],
        fontSize=10,
        leading=14,
        spaceAfter=5,
    )

    story = []

    story.append(
        Paragraph(
            escape(data.logo_text or "LegalEase"),
            title_style,
        )
    )

    story.append(
        Paragraph(
            escape(data.document_type),
            heading_style,
        )
    )

    story.append(
        Paragraph(
            "Key Terms",
            heading_style,
        )
    )

    table_data = [
        ["Term", "Details"]
    ]

    for label, value in build_term_rows(data):
        table_data.append(
            [
                label,
                value or "Not specified"
            ]
        )

    table = Table(
        table_data,
        colWidths=[55 * mm, 115 * mm],
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )

    story.append(table)
    story.append(Spacer(1, 12))

    story.append(
        Paragraph(
            "Generated Document",
            heading_style,
        )
    )

    for line in draft.splitlines():
        if line.strip():
            story.append(
                Paragraph(
                    escape(line),
                    body_style,
                )
            )
        else:
            story.append(
                Spacer(1, 5)
            )

    story.append(Spacer(1, 12))

    story.append(
        Paragraph(
            "<b>LEGAL REVIEW NOTICE:</b> This document is an AI-generated "
            "draft and should be reviewed by a qualified legal professional "
            "before actual use.",
            body_style,
        )
    )

    document.build(story)

    output.seek(0)

    return output
